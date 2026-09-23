"""Regresión de las marcas de Word y su correspondencia con el banco integral."""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX

from import_enfermeria_componente_integral import (
    OUTPUT, SOURCES, answer_marks, build_bank, marked_key,
)


class AnswerMarksTests(unittest.TestCase):
    def options(self):
        document = Document()
        # Los enunciados pueden estar en negrita: solo se pasan las opciones.
        document.add_paragraph().add_run("Enunciado").bold = True
        return document, [document.add_paragraph(value) for value in ("Uno", "Dos", "Tres", "Cuatro")]

    def test_each_supported_mark_selects_its_option(self):
        for name, value in (("bold", True), ("underline", True), ("highlight_color", WD_COLOR_INDEX.YELLOW)):
            with self.subTest(mark=name):
                _, paragraphs = self.options()
                setattr(paragraphs[2].runs[0].font, name, value)
                self.assertEqual(marked_key(paragraphs, "prueba")[0], "C")

    def test_combined_and_split_runs_are_one_answer(self):
        _, paragraphs = self.options()
        paragraphs[1].runs[0].bold = True
        paragraphs[1].add_run(" con continuación").underline = True
        self.assertEqual(marked_key(paragraphs, "prueba"), ("B", ["negrita", "subrayado"]))

    def test_whitespace_and_explicitly_disabled_marks_are_ignored(self):
        _, paragraphs = self.options()
        paragraphs[0].add_run("  ").underline = True
        paragraphs[0].runs[0].underline = False
        paragraphs[0].runs[0].font.highlight_color = WD_COLOR_INDEX.AUTO
        self.assertEqual(answer_marks(paragraphs[0]), set())

    def test_style_inheritance_and_local_override(self):
        document, paragraphs = self.options()
        style = document.styles.add_style("Respuesta", WD_STYLE_TYPE.PARAGRAPH)
        style.font.bold = True
        paragraphs[3].style = style
        self.assertEqual(marked_key(paragraphs, "prueba")[0], "D")
        paragraphs[3].runs[0].bold = False
        self.assertEqual(answer_marks(paragraphs[3]), set())

    def test_no_mark_is_rejected(self):
        _, paragraphs = self.options()
        with self.assertRaisesRegex(ValueError, "encontradas 0"):
            marked_key(paragraphs, "prueba")

    def test_conflicting_marks_are_rejected(self):
        _, paragraphs = self.options()
        paragraphs[0].runs[0].bold = True
        paragraphs[1].runs[0].underline = True
        with self.assertRaisesRegex(ValueError, "encontradas 2"):
            marked_key(paragraphs, "prueba")

    def test_all_source_keys_match_shipped_bank(self):
        questions, audit = build_bank()
        self.assertEqual(questions, json.loads(OUTPUT.read_text()))
        self.assertEqual(len(questions), 43)
        self.assertEqual(len({q["id"] for q in questions}), 43)
        unmarked = [row for row in audit if row["marks"][0].startswith("sin marca")]
        self.assertEqual([row["id"] for row in unmarked], ["local-enfermeria-integral-materno-infantil-010"])
        for question in questions:
            self.assertEqual(len({question[f"option_{key}"] for key in "abcd"}), 4)

    def test_changed_mark_cannot_keep_an_unrelated_explanation(self):
        with tempfile.TemporaryDirectory() as directory:
            source_dir = Path(directory)
            for path in SOURCES.glob("*.docx"):
                shutil.copyfile(path, source_dir / path.name)
            path = source_dir / "PREGUNTAS MATERNO INFANTIL.docx"
            document = Document(path)
            for run in document.paragraphs[3].runs:
                run.font.highlight_color = WD_COLOR_INDEX.AUTO
            document.paragraphs[2].runs[0].underline = True
            document.save(path)
            with self.assertRaisesRegex(ValueError, "falta revisar la explicación"):
                build_bank(source_dir)


if __name__ == "__main__":
    unittest.main()
